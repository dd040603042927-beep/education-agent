from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"E:\education-agent\机器学习期末实践报告_智慧教育智能体平台.docx"


FONT_SONG = "宋体"
FONT_HEI = "黑体"
FONT_EN = "Times New Roman"


def set_run_font(run, name=FONT_SONG, size=14, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_paragraph_format(paragraph, before=0, after=0, line=20, align=None, first_line=True):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = Pt(line)
    if first_line:
        fmt.first_line_indent = Pt(28)
    if align is not None:
        paragraph.alignment = align


def add_text(paragraph, text, font=FONT_SONG, size=14, bold=False, color=None):
    run = paragraph.add_run(text)
    set_run_font(run, font, size, bold, color)
    return run


def add_para(doc, text="", size=14, font=FONT_SONG, bold=False, align=None, before=0, after=0, first_line=True):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, align=align, first_line=first_line)
    add_text(p, text, font=font, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    if level == 1:
        size, before, after = 15, 8, 4
    else:
        size, before, after = 14, 6, 2
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, first_line=False)
    add_text(p, text, font=FONT_HEI, size=size, bold=True)
    return p


def add_center_title(doc, text, size=18):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=8, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_text(p, text, font=FONT_HEI, size=size, bold=True)
    return p


def add_field(paragraph, field_name):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_name
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    set_run_font(run, FONT_HEI, 12, True)


def restart_page_number(section, start=1):
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def setup_section(section, unlink_footer=True):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    if unlink_footer:
        section.footer.is_linked_to_previous = False
        for p in section.footer.paragraphs:
            p.text = ""


def set_body_footer(section):
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = Pt(20)
    add_field(p, "PAGE")


def set_cell_text(cell, text, bold=False, size=11, align=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(16)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, FONT_SONG, size=size, bold=bold)


def set_table_width(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[idx])


def add_table(doc, headers, rows, widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths_cm)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else None
            set_cell_text(cells[i], value, size=10.5, align=align)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(20)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style=None)
    set_paragraph_format(p, before=0, after=0, first_line=False)
    p.paragraph_format.left_indent = Pt(28)
    p.paragraph_format.first_line_indent = Pt(-14)
    add_text(p, "（1）" if False else "• ", font=FONT_SONG, size=14)
    add_text(p, text, font=FONT_SONG, size=14)


doc = Document()
setup_section(doc.sections[0])

# Cover page
add_para(doc, "广东东软学院", size=18, font=FONT_HEI, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=20, first_line=False)
add_para(doc, "实践报告", size=26, font=FONT_HEI, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=90, first_line=False)

cover_fields = [
    ("课程名称", "机器学习"),
    ("项目名称", "智慧教育智能体平台：机器学习知识点问答与学习诊断系统"),
    ("班    级", "24人工1班"),
    ("专    业", "人工智能"),
    ("任课教师", "李卓茜"),
    ("学    号", "24217020127"),
    ("姓    名", "吴永昶"),
]
for label, value in cover_fields:
    p = doc.add_paragraph()
    set_paragraph_format(p, after=3, line=20, first_line=False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(3.2)
    add_text(p, f"{label} ： ", font=FONT_HEI, size=14, bold=True)
    add_text(p, value, font=FONT_SONG, size=14)

add_para(doc, "", after=74, first_line=False)
add_para(doc, "广东东软学院教务部  制", size=12, font=FONT_SONG, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

# Grade page
doc.add_section(WD_SECTION.NEW_PAGE)
setup_section(doc.sections[-1])
add_para(doc, "姓名                            实践报告成绩           ", size=14, font=FONT_HEI, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, first_line=False)
add_para(doc, "评语：", size=14, font=FONT_HEI, bold=True, first_line=False)
for _ in range(14):
    add_para(doc, "", size=14, first_line=False)
add_para(doc, "指导教师（签名）                ", size=14, font=FONT_SONG, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False)
add_para(doc, "年   月   日", size=14, font=FONT_SONG, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False)
add_para(doc, "说明：指导教师评分后，实验报告交院系办公室保存。", size=12, font=FONT_SONG, first_line=False)

# Requirements page
doc.add_section(WD_SECTION.NEW_PAGE)
setup_section(doc.sections[-1])
add_center_title(doc, "实践报告撰写要求", size=18)
requirements = [
    "实践报告按实践项目填写，每个学生做完项目必须撰写并提交实践报告。本科学生实践报告正文不少于3000字，专科学生不少于2500字。",
    "实践报告内容标题由指导老师根据项目实际情况向学生做出具体规定，可参考如下内容：实践目的和任务、实践内容、结论与建议。",
    "实践报告排版要求：",
    "统一用A4（21 X 29.7cm）格式。",
    "标题用小二号黑体加粗，正文用四号宋体。行距为固定值20磅。",
    "页面上边距2.54cm，下边距2.54 cm，左边距3cm，右边距2.2cm。",
    "实践报告页码从正文页面起计算。页码字号，选用小四号粗黑体并居中。",
]
for item in requirements:
    add_para(doc, item, size=14, font=FONT_SONG)

# Body section with page numbers
body_section = doc.add_section(WD_SECTION.NEW_PAGE)
setup_section(body_section)
restart_page_number(body_section, 1)
set_body_footer(body_section)

add_center_title(doc, "正文", size=18)
add_heading(doc, "一、实践目的和任务", 1)
purpose_paragraphs = [
    "本实践项目以当前项目“智慧教育智能体平台”为基础，围绕机器学习课程中“学习诊断、知识表示、个性化反馈、算法实践验证”四类核心问题进行系统化实现与分析。项目不是单一算法的离线演示，而是把机器学习课程知识点、学生学习行为、教师教学活动和大模型工作流整合为一个可运行的教育智能体平台，使机器学习知识能够在真实教学场景中被检索、诊断、反馈和验证。",
    "本项目选择教育智能体作为实践对象，原因在于机器学习课程本身具有概念密集、算法流程多、评价指标复杂、学生易混淆点集中的特点。传统实验往往只关注某个算法的训练结果，而本项目进一步关注学习者是否理解算法的任务类型、输入输出、模型假设、目标函数、评价指标和失误原因。平台通过知识图谱、课程资料 RAG、Dify 多节点诊断工作流和模型代码实验室，把这些学习维度落到可交互的产品功能中。",
    "实践的总体目标是完成一个本地可运行、前后端闭环、能够支撑教师和学生两种角色的机器学习辅助教学系统。学生端可以围绕知识点进行问答、诊断、练习、作业提交和学习路径查看；教师端可以上传课程资料、生成知识图谱、管理班级和作业、调用 AI 助教完成备课、出题、批改和学情分析。后端负责身份认证、角色权限、课程资料解析、图谱任务、学习事件记录、掌握度更新、Dify 回调同步和代码执行安全控制。",
    "本实践的具体任务包括：第一，梳理机器学习课程知识结构，构建可用于 RAG 检索和学习诊断的知识库与知识图谱；第二，设计学生学习诊断流程，使系统能够根据问题、学生回答、资料证据和知识点定位输出结构化反馈；第三，实现教师与学生双端交互，覆盖课程资料、知识图谱、作业、聊天和个人学习画像；第四，建设机器学习算法实验室，让算法代码能够被真实运行、返回 stdout、stderr、退出码和耗时，而不是展示固定示例；第五，完成安全、存储和可发布基线，保证项目具备持续迭代和小范围内测条件。",
]
for text in purpose_paragraphs:
    add_para(doc, text)

add_heading(doc, "二、实践内容", 1)
add_heading(doc, "（一）项目总体结构与技术路线", 2)
overview = [
    "当前项目是一个不依赖外网安装的本地全栈原型。后端入口为 server.js，使用 Node.js 原生 http 模块、文件系统、加密、压缩和子进程能力实现 API 服务；前端位于 public/index.html、public/styles.css 和 public/app.js，是一个静态 SPA；数据默认保存到 data/db.json，同时提供 SQLite 迁移与镜像存储能力。项目启动后默认访问地址为 http://127.0.0.1:5107，也支持通过环境变量指定端口、数据目录、日志目录、Dify 工作流地址和安全密钥。",
    "系统技术路线可以概括为“课程资料解析 -> 文本分块与本地语义指纹 -> 知识图谱与 GraphRAG 上下文 -> 学习诊断工作流 -> 学习事件与掌握度数据库 -> 前端角色化反馈”。课程资料进入系统后被解析为可检索文本片段，后端使用关键词和 64 维本地语义指纹进行召回和重排；知识图谱节点补充前置依赖、易混淆关系、考察属性和学习路径信息；Dify 工作流负责更复杂的知识点定位、证据整理、掌握度评分和诊断反馈；回调结果再写入学习画像、错题和会话记录。",
]
for text in overview:
    add_para(doc, text)

add_table(
    doc,
    ["维度", "项目事实", "实践意义"],
    [
        ["运行方式", "本地 Node.js 后端 + 静态 SPA 前端，默认端口 5107。", "便于课程验收、离线演示和局域网小范围内测。"],
        ["知识工程", "资料分块参数为 900 字符、160 字符重叠，语义指纹维度为 64。", "在不依赖外部向量库的情况下完成课程资料召回和重排。"],
        ["学习诊断", "学生端对接 Dify 学习诊断工作流，教师端对接独立教学工作流。", "区分学生问答诊断与教师备课、批改、学情分析任务。"],
        ["数据闭环", "learning_events、diagnosis_results、student_mastery 构成真实学习数据落库结构。", "把 AI 反馈转化为可追踪的学习事件和掌握度快照。"],
        ["算法实践", "模型实验室通过 Python 子进程真实运行代码并返回退出码、输出和耗时。", "避免只展示静态答案，体现机器学习实验的可验证性。"],
        ["安全基线", "HttpOnly session、scrypt 密码哈希、RBAC、审计日志、上传白名单。", "满足教学平台发布和小范围内测的基础要求。"],
    ],
    [2.2, 6.2, 6.4],
)

add_heading(doc, "（二）需求分析与角色设计", 2)
need_text = [
    "从机器学习课程实践角度看，系统至少需要解决三类需求。第一是学生学习需求：学生需要围绕 KNN、线性回归、过拟合与泛化、逻辑回归、矩阵分解、神经网络、CNN、RNN、SVM、决策树、集成学习、KMeans、PCA、朴素贝叶斯、EM/GMM、自编码器等知识点进行提问、练习和诊断；系统不仅要回答概念，还要指出任务类型、输入输出、目标函数、评价指标和易错点。第二是教师教学需求：教师需要把教材、讲义和补充资料上传为课程资料，生成知识图谱，基于班级、资料、图谱、作业和提交记录调用 AI 助教完成备课、出题、批改与学情分析。第三是平台运维需求：系统需要保存账号、班级、资料、图谱、作业、会话、错题、掌握度和审计日志，并在服务重启后保留上传会话和图谱生成任务状态。",
    "项目据此设计了管理员、教师和学生三种角色。管理员侧重数据导出和系统管理；教师端包含 AI 助教、作业管理、知识图谱、课程资料、模型实验室、班级管理、师生消息和个人信息；学生端包含学习周期驾驶舱、AI 诊断与辅导、个人学习路径、学习证据档案、学习周期申报、课程、节点练习/作业、学习作品/实验、站内消息和个人信息。角色分离使同一个 AI 入口能够根据身份加载不同上下文，避免把学生学习诊断逻辑误用于教师批改或备课场景。",
]
for text in need_text:
    add_para(doc, text)

add_heading(doc, "（三）课程资料 RAG 与知识图谱构建", 2)
rag_text = [
    "项目中的课程资料处理体现了机器学习系统的典型数据工程流程。教师可上传 PDF、TXT、Markdown、Office 文档、图片等格式，后端先进行文件格式白名单和基础文件头校验，再将资料解析为文本。对于 PDF，项目优先调用 Python 智能体使用 PyMuPDF 读取文本层，失败时尝试 pdfplumber 和 pypdf；如果文本层不足，则用 PyMuPDF 渲染页面并调用 PaddleOCR 对扫描版页面做有限 OCR，最后才回退到 Node 轻量解析器。这样的多级策略兼顾了普通电子教材和扫描版资料，符合真实教学资料来源复杂的事实。",
    "资料文本进入 RAG 流程后，系统按段落和句末标点进行切分，单片段约 900 字符，片段之间保留 160 字符重叠，并为每个片段生成关键词和 64 维本地语义指纹。后端检索时同时考虑关键词匹配和语义相似度，再将高分片段作为课程资料证据传给 AI 工作流。由于本项目定位为本地原型，这种轻量检索方案降低了部署成本；同时，项目说明中也明确后续可以升级为真实向量数据库和更完整的重排模型。",
    "知识图谱是本项目区别于普通问答系统的关键模块。后端把课程内容组织为学科知识图谱、认知能力图谱、教学资源图谱、素养与价值图谱、学习者认知图谱等层次，并为节点增强 GraphRAG 信息。图谱关系包括 contains、prerequisite、misconception、cross-link、assessment、resource、competency 和 semantic。每个节点除名称和层级外，还带有前置依赖、易混淆提醒、考察属性、学习路径角色和局部子图检索提示。学生双击节点可以查看知识路径、前置依赖、教学资源、考察属性和详情，教师也可以导入、导出、删除和上传总图谱。",
    "从机器学习角度看，知识图谱在系统中承担了特征组织与结构化先验的作用。RAG 负责从非结构化资料中召回文本证据，图谱负责给出层级、依赖和易错关系。二者结合后，系统能够在回答“什么是 KNN”时不仅给出定义，还能提示距离度量、标准化、K 值选择、分类/回归评价指标等关联知识；在学生答错时，也可以沿 prerequisite 关系追溯真正薄弱的前置节点。",
]
for text in rag_text:
    add_para(doc, text)

add_heading(doc, "（四）机器学习学习诊断工作流设计", 2)
workflow_text = [
    "学生端 AI 助教保留统一入口 /api/ai/chat，但后端按角色拆分工作流。学生端接入 dify/ml_learning_diagnosis/ml_learning_diagnosis_assistant_upgraded_0_6_0.yml，教师端接入独立的教师教学工作流。学生请求会自动携带 student_id、学科、章节、知识点、问题、学生答案、课程资料上下文、错因上下文、掌握度上下文、图谱上下文和回调地址；教师请求则携带 teacher_id、task_type、class_id、selected_material_ids、selected_graph_id、homework_id、student_submission_id、rubric、project_rag_context 和 project_class_context 等信息。",
    "学习诊断工作流按节点完成任务拆分：开始节点接收问题和上下文；输入清洗与任务识别节点规范化问题并判断问答或诊断模式；课程知识库 RAG 节点检索教师端课程资料；错因知识库 RAG 节点检索易混淆点、题库和评分标准；RAG 片段去重与证据整理节点合并可追溯证据；朴素贝叶斯知识点分类节点定位知识点；掌握度评分节点计算学习画像；LLM 标准解释节点生成可学习的标准答案；条件分支区分问答或诊断；LLM 诊断反馈节点给出错因和修正建议；结构化 JSON 输出节点返回 final_answer、citations、warnings、follow_up_actions 等字段；项目数据库同步 HTTP 节点把结构化结果回写平台。",
    "诊断结果不是停留在聊天文本中，而是进入学习数据闭环。后端 /api/integrations/dify/diagnosis-callback 通过 DIFY_CALLBACK_TOKEN 校验回调来源，解析 callback_payload 后同步会话、错题、学习画像和诊断结果。learning_events 记录实际发生的学习行为，diagnosis_results 保存 AI 或教师评估的结构化输出，student_mastery 保存学生在某学科某知识点上的当前掌握度。项目还提供 /api/mastery/update 接口，允许在用户标记掌握、知识测试、教师确认批改等场景中更新掌握度。",
    "这种设计体现了机器学习实践中“模型输出必须可追踪、可评价、可迭代”的原则。系统不只给出自然语言建议，还保存 topic_label、mastery_score、mastery_level、error_tags、missing_points、rag_evidence 和 final_answer。后续教师可以基于这些结构化字段进行班级学情分析，学生也可以在个人学习路径中看到薄弱点和复习建议。",
]
for text in workflow_text:
    add_para(doc, text)

add_heading(doc, "（五）模型实验室与算法代码验证", 2)
lab_text = [
    "模型实验室是本项目最直接体现机器学习实践能力的模块。前端提供画布代码编辑器和模型组件，后端提供 /api/model-code/generate 与 /api/model-code/run。生成接口会根据用户需求、学科和课程资料召回结果生成 Python 教学实验代码；运行接口会把代码写入 data/runtime 下的临时目录，通过 Python 子进程真实执行，并返回 success、exitCode、timedOut、durationMs、stdout、stderr、pythonCommand 和 output。README 中也明确指出，机器学习算法实验室返回的是真实 stdout/stderr/退出码/耗时，而不是固定展示预设结果。",
    "为保证教学代码执行安全，后端对模型代码做了静态安全检查。默认只允许 math、random、statistics、collections、itertools、functools、operator、heapq、bisect、decimal、fractions、typing、dataclasses 等标准库导入；同时拦截 open、eval、exec、compile、input、__import__、globals、locals、getattr、setattr、system、popen、spawn、remove、unlink 等危险能力，以及 Python 内部对象访问特征。执行超时时间默认限制在 10 秒，输出长度和代码长度也有上限。",
    "项目还实现了算法生成后的验证与自动修复机制。verifyGeneratedAlgorithm 会先运行生成的代码，若成功且 stdout 非空，则记录已在服务器真实执行验证；如果运行失败且配置了 OpenAI API Key，runModelCodeWithAutoRepair 可以结合课程资料上下文、错误输出和原始需求进行有限次数修复，再次运行直至成功或达到上限。实验记录会保存 agentName、subject、prompt、codeMode、difficulty、sourceType、title、chapter、citations、explanation、verifiedRun、repairAttempts、repairHistory 和 workflow。这个流程使学生能够看到“课程资料检索、代码生成、Python 沙箱运行、错误修复、实验记录保存”的完整链路。",
]
for text in lab_text:
    add_para(doc, text)

add_heading(doc, "（六）数据持久化、安全与发布基线", 2)
security_text = [
    "项目默认使用 JSON 文件作为存储，首次启动会生成演示数据；同时提供 SQLite 迁移能力。storage/schema.sql 中定义了 learning_events、diagnosis_results、student_mastery、external_id_mappings 和 ingestion_jobs 等表，docs/learning-data-database.md 说明 SQLite 适配器会把当前 app_state 保持为兼容结构，并把真实学习数据镜像到规范化表。这是一种渐进式数据库改造：既不破坏现有 /api/state 返回结构，又为后续 PostgreSQL、任务队列、真实向量检索和多租户学校模型预留路径。",
    "安全方面，项目已经补齐小范围内测所需的第一批基线。登录态使用后端签名的 HttpOnly session cookie，不再信任前端伪造的 userId；SESSION_SECRET 支持强随机配置，COOKIE_SECURE 可在 HTTPS 反向代理场景启用；密码使用 Node 内置 scrypt 哈希，旧明文密码会自动迁移；后端执行基础 RBAC，支持 admin、teacher、student 角色；登录失败次数限制和冷却保护降低暴力破解风险；敏感操作写入 data/db.json 的 auditLogs，并追加到 logs/audit.log。",
    "上传和运行也做了防护。课程资料上传使用格式白名单和文件头校验，阻断可执行文件、脚本类文件和伪装格式；PDF 解析设置最大文件大小、流数量、压缩流和解压输出限制；图谱生成和上传会话持久化到 data/runtime，服务重启后不会直接丢失任务状态；data/db.json 写入采用临时文件加 rename 的原子写入方式。虽然这仍然是本地原型，但这些设计已经体现出从课堂实验到可发布系统的工程意识。",
]
for text in security_text:
    add_para(doc, text)

add_heading(doc, "（七）测试与验证结果", 2)
verify_intro = "为了保证报告内容基于项目真实状态，本次在当前工作区执行了项目自带检查。PowerShell 默认 npm.ps1 因执行策略被拦截后，改用 npm.cmd 执行脚本。检查时间为 2026 年 6 月 25 日，结果如下。"
add_para(doc, verify_intro)
add_table(
    doc,
    ["验证项", "命令", "结果", "说明"],
    [
        ["语法检查", "npm.cmd run check", "通过", "依次对 server.js、public/app.js、storage 模块和脚本文件执行 node --check。"],
        ["发布烟测", "npm.cmd test", "通过", "scripts/release_smoke_test.js 输出 release smoke test passed。"],
        ["测试覆盖重点", "release_smoke_test.js", "覆盖", "健康检查、首页资源、登录 cookie、会话恢复、越权阻断、恶意上传拦截、Dify 工作流模拟、图谱任务、机器学习代码真实执行和 AI 对话结构。"],
    ],
    [2.3, 4.0, 2.0, 5.7],
)
test_text = [
    "从测试脚本内容看，烟测并不是简单启动检查，而是创建临时数据目录、启动 mock Dify 服务、启动项目后端、访问健康检查和首页资源、验证静态资源缓存策略、检查前端 bundle 中角色工作台和 CSP 安全表达式绘图逻辑、模拟上传课程资料、等待图谱任务完成，并验证 AI 对话结构。测试还覆盖了登录 cookie、会话恢复、权限阻断、恶意上传拦截和模型代码运行等关键路径，能够较好支撑“项目可运行、关键功能闭环成立”的结论。",
    "测试结果说明当前项目至少满足课程实践验收中的可运行性和基本可靠性要求。与只提交静态页面或算法片段相比，本项目具备完整的前后端交互、数据持久化、AI 工作流对接、代码执行和验证脚本，因此实践成果更接近真实软件工程项目。",
]
for text in test_text:
    add_para(doc, text)

add_heading(doc, "（八）项目创新点与课程知识对应关系", 2)
innovation_rows = [
    ["知识图谱 + RAG", "把机器学习课程知识点、前置依赖、易混淆点和资料证据结合，提升问答可追溯性。", "知识表示、信息检索、语义相似度、图结构建模。"],
    ["结构化学习诊断", "输出 topic_label、mastery_score、error_tags、missing_points、rag_evidence 等字段。", "分类、评分、特征抽取、模型评价、错误分析。"],
    ["模型代码实验室", "代码真实运行并返回 stdout、stderr、退出码和耗时，支持生成后验证与修复。", "实验可复现、算法实现、运行时验证、调参反馈。"],
    ["学习数据闭环", "学习事件、诊断结果、掌握度快照形成可分析的学生画像。", "数据建模、监督信号、时间序列学习记录、个性化推荐。"],
    ["安全与发布基线", "RBAC、scrypt、HttpOnly cookie、审计日志、上传校验、原子写入。", "机器学习系统工程化、数据安全、平台可靠性。"],
]
add_table(doc, ["创新点", "项目实现", "对应机器学习能力"], innovation_rows, [3.0, 6.3, 4.7])

innovation_text = [
    "本项目的主要创新不在于复现单个经典算法，而在于把机器学习知识应用于教育智能体的完整闭环。知识图谱负责提供结构化先验，RAG 负责提供可追溯证据，Dify 工作流负责复杂诊断推理，学习事件数据库负责结果沉淀，模型实验室负责算法代码的真实执行验证。多个模块共同支撑“学、练、诊、评、改”的过程，使机器学习课程从静态知识点转化为动态学习系统。",
    "项目还体现了较强的工程完整性。前端不是单页展示，而是区分教师、学生、管理员角色；后端不是简单代理，而是包含上传、解析、检索、图谱任务、权限、审计、存储迁移、Dify 回调、代码沙箱和测试脚本；数据层不是临时变量，而是从 JSON 逐步过渡到 SQLite 规范表。这些事实使项目可以作为机器学习期末实践报告的真实支撑材料。",
]
for text in innovation_text:
    add_para(doc, text)

add_heading(doc, "三、结论与建议", 1)
conclusion_text = [
    "通过本次实践，项目已经完成一个面向机器学习课程的智慧教育智能体平台原型。系统能够支持教师上传课程资料并生成知识图谱，支持学生围绕机器学习知识点进行 AI 问答、诊断、练习和作业提交，支持教师进行备课、出题、批改和学情分析，支持学习事件和掌握度数据的沉淀，也支持机器学习算法代码的真实运行验证。项目检查和发布烟测均已通过，说明当前实现具备基本可运行性和较好的工程闭环。",
    "从课程目标看，本项目覆盖了机器学习实践中常见的关键能力：数据采集与预处理、知识表示、文本检索、图结构建模、分类与评分、评价指标、错误分析、代码实验、结果验证和系统工程化。尤其是学习诊断工作流把机器学习知识点定位、RAG 证据、掌握度评分、错因标签和后续建议统一到结构化输出中，使系统能够服务真实教学，而不是停留在单轮问答。",
    "项目当前仍有改进空间。第一，RAG 目前采用本地轻量语义指纹，适合本地原型，但在更大规模资料和多班级使用时，应升级为真实向量检索、rerank 模型和缓存策略。第二，学习掌握度更新目前主要依赖 Dify 回调、知识测试、教师确认批改和用户标记，后续可引入更严谨的知识追踪模型，例如 Bayesian Knowledge Tracing、Deep Knowledge Tracing 或基于图结构的掌握度传播。第三，模型实验室目前以标准库 Python 教学代码为主，后续可在安全沙箱成熟后增加 numpy、pandas、scikit-learn 等受控依赖，并提供数据集管理和实验版本对比。第四，当前项目仍建议在长期运行前迁移到 SQLite WAL 或 PostgreSQL，引入任务队列、备份恢复、监控和端到端测试。",
    "总体而言，本实践项目能够把机器学习课程知识、学习者诊断、知识图谱、RAG、大模型工作流和软件工程实现结合起来，既体现了机器学习理论的应用价值，也体现了完整系统开发能力。若后续继续完善向量检索、知识追踪模型、实验数据集和教师评价闭环，该平台可以从课程期末实践进一步发展为可用于真实课堂的小型智能教学系统。",
]
for text in conclusion_text:
    add_para(doc, text)

add_heading(doc, "参考资料与项目依据", 1)
refs = [
    "E:\\education-agent\\README.md：项目启动、Dify 对接、发行基线、已实现模块说明。",
    "E:\\education-agent\\server.js：后端 API、RAG 分块、知识图谱增强、Dify 回调、模型代码运行、安全与存储逻辑。",
    "E:\\education-agent\\public\\app.js：教师端、学生端、管理员端前端 SPA 与交互模块。",
    "E:\\education-agent\\storage\\schema.sql：learning_events、diagnosis_results、student_mastery 等学习数据表结构。",
    "E:\\education-agent\\docs\\learning-data-database.md：学习数据数据库集成路线与迁移说明。",
    "E:\\education-agent\\dify\\ml_learning_diagnosis\\ml_learning_diagnosis_assistant_upgraded_0_6_0.yml：机器学习诊断 Dify 工作流。",
    "E:\\education-agent\\scripts\\release_smoke_test.js：发布烟测与关键功能验证脚本。",
]
for ref in refs:
    add_para(doc, ref, first_line=False)

# Final paragraph/style audit defaults.
for section in doc.sections:
    setup_section(section, unlink_footer=False)

doc.save(OUT)
print(OUT)
