import zipfile

from docx import Document


DOCX = r"E:\education-agent\机器学习期末实践报告_智慧教育智能体平台.docx"


doc = Document(DOCX)
text = "\n".join(p.text for p in doc.paragraphs)
print("paragraphs", len(doc.paragraphs))
print("tables", len(doc.tables))
print("sections", len(doc.sections))
print("chars_no_space", len("".join(text.split())))
print("identity_fields", all(item in text for item in ["吴永昶", "24217020127", "24人工1班", "李卓茜"]))
for i, section in enumerate(doc.sections):
    print(
        "section",
        i,
        round(section.page_width.cm, 2),
        round(section.page_height.cm, 2),
        round(section.top_margin.cm, 2),
        round(section.bottom_margin.cm, 2),
        round(section.left_margin.cm, 2),
        round(section.right_margin.cm, 2),
        [p.text for p in section.footer.paragraphs],
    )
for i, table in enumerate(doc.tables):
    print("table", i, len(table.rows), len(table.columns), [cell.text for cell in table.rows[0].cells])

with zipfile.ZipFile(DOCX) as zf:
    document_xml = zf.read("word/document.xml").decode("utf-8")
    print("pgNumType", "pgNumType" in document_xml)
    print("page_start_1", 'w:start="1"' in document_xml)
    print("fixed_line_20pt_count", document_xml.count('w:line="400"'))
    print("song_font_count", document_xml.count("宋体"))
    print("hei_font_count", document_xml.count("黑体"))
    for name in zf.namelist():
        if name.startswith("word/footer"):
            footer_xml = zf.read(name).decode("utf-8")
            print("footer", name, "PAGE" in footer_xml, footer_xml.count("fldChar"))
